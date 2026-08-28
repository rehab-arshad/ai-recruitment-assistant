import os
os.environ["HF_HUB_OFFLINE"] = "1"

import re
import hashlib
import chromadb
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
DB_PATH = "chroma_store"
COLLECTION_NAME = "candidates"

model = SentenceTransformer(MODEL_NAME)


def n_tokens(text):
    return max(1, len(text) // 4)


def load_pdf(path):
    reader = PdfReader(path)
    pages = []

    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
        text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
        text = re.sub(r"[ \t]{2,}", " ", text)
        text = text.strip()

        if text:
            pages.append({"page": i, "text": text})

    if not pages:
        raise ValueError("No text found. This PDF is probably a scan.")
    return pages


def load_docx(path):
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # tables are separate from paragraphs in python-docx — many resume
    # templates put the whole skills section in one
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("No text found in this document.")

    return [{"page": 1, "text": text}]


def load_txt(path):
    with open(path, encoding="utf-8", errors="replace") as fh:
        text = fh.read().strip()

    if not text:
        raise ValueError("File is empty.")

    return [{"page": 1, "text": text}]


def load_document(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".txt", ".md"}:
        return load_txt(path)
    if ext == ".doc":
        raise ValueError("Legacy .doc not supported — save as .docx and retry.")

    raise ValueError(f"Unsupported file type: {ext}")


def split_paragraph(para, max_tokens):
    sentences = re.split(r"(?<=[.!?])\s+", para)
    parts, buf = [], ""

    for sent in sentences:
        candidate = f"{buf} {sent}".strip()
        if buf and n_tokens(candidate) > max_tokens:
            parts.append(buf)
            buf = sent
        else:
            buf = candidate

    if buf:
        parts.append(buf)
    return parts


def chunk_pages(pages, chunk_tokens=250, overlap_tokens=50):
    block_cap = chunk_tokens - overlap_tokens

    blocks = []
    for p in pages:
        for para in p["text"].split("\n\n"):
            para = para.strip()
            if not para:
                continue
            if n_tokens(para) <= block_cap:
                blocks.append((p["page"], para, n_tokens(para)))
            else:
                for part in split_paragraph(para, block_cap):
                    blocks.append((p["page"], part, n_tokens(part)))

    def build(window):
        return {
            "text": "\n\n".join(b[1] for b in window),
            "page_start": window[0][0],
            "page_end": window[-1][0],
            "n_tokens": sum(b[2] for b in window),
        }

    chunks, window, total = [], [], 0

    for block in blocks:
        if window and total + block[2] > chunk_tokens:
            chunks.append(build(window))

            tail, tail_total = [], 0
            for b in reversed(window):
                if tail_total + b[2] > overlap_tokens:
                    break
                tail.insert(0, b)
                tail_total += b[2]

            window, total = tail, tail_total

        window.append(block)
        total += block[2]

    if window:
        chunks.append(build(window))

    return chunks


def embed_texts(texts, show_progress=False):
    return model.encode(
        texts,
        normalize_embeddings=True,
        show_progress_bar=show_progress,
    )


def get_collection():
    chroma = chromadb.PersistentClient(path=DB_PATH)
    return chroma.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )


def index_pdf(path, force=False):
    collection = get_collection()
    source = os.path.basename(path)

    existing = collection.get(where={"source": source})
    if existing["ids"] and not force:
        print(f"'{source}' already indexed — {len(existing['ids'])} chunks.")
        return collection, source

    if existing["ids"]:
        collection.delete(ids=existing["ids"])

    pages = load_document(path)
    chunks = chunk_pages(pages)
    print(f"{source}: {len(pages)} page(s) -> {len(chunks)} chunks")

    vectors = embed_texts([c["text"] for c in chunks], show_progress=True)

    ids = [
        hashlib.sha1(f"{source}:{i}:{c['text']}".encode()).hexdigest()[:16]
        for i, c in enumerate(chunks)
    ]

    collection.add(
        ids=ids,
        documents=[c["text"] for c in chunks],
        embeddings=[v.tolist() for v in vectors],
        metadatas=[
            {"source": source, "page_start": c["page_start"], "page_end": c["page_end"]}
            for c in chunks
        ],
    )

    print(f"indexed — collection holds {collection.count()} chunks")
    return collection, source